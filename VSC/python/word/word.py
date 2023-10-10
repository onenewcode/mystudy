from docx import Document
from docx.shared import Inches,Cm,Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT #导入的关于word操作的相关模块
import datetime #获取当前日期

docx = Document(r"C:\Users\ztf\OneDrive\桌面\艾瑞森2022年新增.docx")#docx文件的地址
tables = docx.tables#获取所有表格
table=tables[0]#获取第一个表格
tablerow=len(table.rows)#获取行的总数
tablecol=len(table.columns)#获取列的总数
# for i in range(tablerow):
#     print(table.rows[i].cells)
# 遍历
for i in range(tablerow):
    for j in range(tablerow):
        print(table.cell(i,j).text)

